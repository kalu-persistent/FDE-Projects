import hashlib
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pinecone import Pinecone, ServerlessSpec
from sqlalchemy import select

from .config import get_settings
from .database import SessionLocal, init_db
from .models import IngestedDocument
from .services import embeddings


settings = get_settings()

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

FOLDER_PERMISSIONS = {
    "public": ["ALL"],
    "hr-restricted": ["HR", "ADMIN"],
}


def read_document(file_path: Path) -> str:
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        reader = PdfReader(file_path)

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if extension == ".docx":
        document = Document(file_path)

        content = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        # Also read tables from DOCX files.
        for table in document.tables:
            for row in table.rows:
                content.append(
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                )

        return "\n".join(content)

    if extension == ".txt":
        return file_path.read_text(
            encoding="utf-8",
            errors="ignore"
        )

    raise ValueError(
        f"Unsupported document type: {extension}"
    )


def normalize_text(text: str) -> str:
    return " ".join(text.split())


def create_chunks(
    text: str,
    chunk_size: int = 900,
    chunk_overlap: int = 120
) -> list[str]:

    step = chunk_size - chunk_overlap

    return [
        text[index:index + chunk_size]
        for index in range(0, len(text), step)
        if text[index:index + chunk_size].strip()
    ]


def calculate_file_hash(file_path: Path) -> str:
    hash_object = hashlib.sha256()

    with file_path.open("rb") as file:

        for block in iter(
            lambda: file.read(1024 * 1024),
            b""
        ):
            hash_object.update(block)

    return hash_object.hexdigest()


def get_allowed_roles(
    file_path: Path,
    documents_folder: Path
) -> list[str]:

    relative_parts = [
        part.lower()
        for part in file_path
        .relative_to(documents_folder)
        .parts[:-1]
    ]

    for folder_name, roles in FOLDER_PERMISSIONS.items():

        if folder_name in relative_parts:
            return roles

    # Secure default if document is not inside
    # public or hr-restricted.
    return ["ADMIN"]


def get_pinecone_index(embedding_dimension: int):
    pinecone = Pinecone(
        api_key=settings.pinecone_api_key
    )

    index_names = pinecone.list_indexes().names()

    if settings.pinecone_index not in index_names:

        pinecone.create_index(
            name=settings.pinecone_index,
            dimension=embedding_dimension,
            metric="cosine",
            spec=ServerlessSpec(
                cloud="aws",
                region="us-east-1"
            )
        )

    return pinecone.Index(
        settings.pinecone_index
    )


def ingest_all_documents(
    documents_folder: str = "data/documents"
) -> dict:

    if not settings.pinecone_api_key:
        raise RuntimeError(
            "PINECONE_API_KEY is not configured"
        )

    root_folder = Path(documents_folder).resolve()

    if not root_folder.exists():
        raise RuntimeError(
            f"Document folder not found: {root_folder}"
        )

    files = sorted([
        file
        for file in root_folder.rglob("*")
        if file.is_file()
        and file.suffix.lower() in SUPPORTED_EXTENSIONS
    ])

    embedding_model = embeddings()

    vector_index = get_pinecone_index(
        embedding_model
        .get_sentence_embedding_dimension()
    )

    result = {
        "found": len(files),
        "indexed": 0,
        "skipped": 0,
        "failed": 0,
        "documents": []
    }

    for file_path in files:

        relative_path = (
            file_path
            .relative_to(root_folder)
            .as_posix()
        )

        content_hash = calculate_file_hash(
            file_path
        )

        allowed_roles = get_allowed_roles(
            file_path,
            root_folder
        )

        acl_key = ",".join(
            sorted(allowed_roles)
        )

        try:
            # First check PostgreSQL before creating
            # chunks or embeddings.
            with SessionLocal() as database:

                existing_document = database.scalar(
                    select(IngestedDocument).where(
                        IngestedDocument.content_hash
                        == content_hash,

                        IngestedDocument.acl_key
                        == acl_key,

                        IngestedDocument.status
                        == "COMPLETED"
                    )
                )

                if existing_document:

                    result["skipped"] += 1

                    result["documents"].append({
                        "source": relative_path,
                        "status": "SKIPPED_ALREADY_EXISTS",
                        "chunks":
                            existing_document.chunk_count,
                        "roles": allowed_roles
                    })

                    continue

            extracted_text = read_document(
                file_path
            )

            normalized_text = normalize_text(
                extracted_text
            )

            if not normalized_text:
                raise RuntimeError(
                    "No extractable text found"
                )

            chunks = create_chunks(
                normalized_text
            )

            vectors = embedding_model.encode(
                chunks,
                normalize_embeddings=True
            ).tolist()

            pinecone_records = []

            for chunk_number, (
                chunk,
                vector
            ) in enumerate(
                zip(chunks, vectors)
            ):

                # Stable vector ID prevents duplicates.
                vector_id = hashlib.sha256(
                    (
                        f"{content_hash}:"
                        f"{acl_key}:"
                        f"{chunk_number}"
                    ).encode()
                ).hexdigest()

                pinecone_records.append({
                    "id": vector_id,
                    "values": vector,
                    "metadata": {
                        "text": chunk,
                        "source": relative_path,
                        "filename": file_path.name,
                        "content_hash": content_hash,
                        "chunk": chunk_number,
                        "allowed_roles": allowed_roles
                    }
                })

            # Upload in batches.
            for start in range(
                0,
                len(pinecone_records),
                50
            ):

                vector_index.upsert(
                    vectors=pinecone_records[
                        start:start + 50
                    ],
                    namespace=settings.pinecone_namespace
                )

            # Save ingestion status in PostgreSQL.
            with SessionLocal() as database:

                registry_record = IngestedDocument(
                    source_path=relative_path,
                    content_hash=content_hash,
                    acl_key=acl_key,
                    chunk_count=len(pinecone_records),
                    status="COMPLETED"
                )

                database.add(registry_record)
                database.commit()

            result["indexed"] += 1

            result["documents"].append({
                "source": relative_path,
                "status": "INDEXED",
                "chunks": len(pinecone_records),
                "roles": allowed_roles
            })

        except Exception as exception:

            result["failed"] += 1

            result["documents"].append({
                "source": relative_path,
                "status": "FAILED",
                "error": str(exception),
                "roles": allowed_roles
            })

    return result


if __name__ == "__main__":
    init_db()

    response = ingest_all_documents()

    print(
        json.dumps(
            response,
            indent=2,
            default=str
        )
    )